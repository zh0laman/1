"""
Парсеры документов — извлечение чистого текста из разных форматов.
Интегрировано из rag.zip в основной проект alem-ai-backend.
"""
import io
import logging
import mimetypes
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional

logger = logging.getLogger(__name__)


@dataclass
class ParsedPage:
    page_number: int
    text: str
    metadata: dict = field(default_factory=dict)


@dataclass
class ParsedDocument:
    pages: List[ParsedPage]
    title: Optional[str] = None
    author: Optional[str] = None
    metadata: dict = field(default_factory=dict)

    @property
    def full_text(self) -> str:
        return "\n\n".join(p.text for p in self.pages if p.text.strip())


def parse_pdf(data: bytes) -> ParsedDocument:
    try:
        import fitz
        doc = fitz.open(stream=data, filetype="pdf")
        pages: List[ParsedPage] = []
        meta = doc.metadata or {}
        title = meta.get("title") or None
        author = meta.get("author") or None
        for page_num, page in enumerate(doc, start=1):
            text = page.get_text("text")
            if text.strip():
                pages.append(ParsedPage(page_number=page_num, text=text.strip(), metadata={"page": page_num}))
        doc.close()
        if not pages:
            return _parse_pdf_fallback(data)
        return ParsedDocument(pages=pages, title=title, author=author)
    except Exception as e:
        logger.warning("PyMuPDF failed, trying fallback: %s", e)
        return _parse_pdf_fallback(data)


def _parse_pdf_fallback(data: bytes) -> ParsedDocument:
    try:
        import PyPDF2
        reader = PyPDF2.PdfReader(io.BytesIO(data))
        pages: List[ParsedPage] = []
        meta = reader.metadata or {}
        for i, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                pages.append(ParsedPage(page_number=i, text=text.strip()))
        return ParsedDocument(pages=pages, title=meta.get("/Title"), author=meta.get("/Author"))
    except Exception as e:
        logger.error("PDF parse completely failed: %s", e)
        return ParsedDocument(pages=[])


def parse_docx(data: bytes) -> ParsedDocument:
    try:
        from docx import Document
        from docx.oxml.ns import qn
        doc = Document(io.BytesIO(data))
        paragraphs: List[str] = []
        for element in doc.element.body:
            tag = element.tag.split("}")[-1]
            if tag == "p":
                text = "".join(run.text for run in element.iter(qn("w:t")))
                if text.strip():
                    paragraphs.append(text.strip())
            elif tag == "tbl":
                for row in element.iter(qn("w:tr")):
                    cells = []
                    for cell in row.iter(qn("w:tc")):
                        cell_text = "".join(t.text for t in cell.iter(qn("w:t")))
                        cells.append(cell_text.strip())
                    if any(cells):
                        paragraphs.append(" | ".join(cells))
        full_text = "\n".join(paragraphs)
        pages = [ParsedPage(page_number=1, text=full_text)] if full_text.strip() else []
        props = doc.core_properties
        return ParsedDocument(pages=pages, title=props.title or None, author=props.author or None)
    except Exception as e:
        logger.error("DOCX parse failed: %s", e)
        return ParsedDocument(pages=[])


def parse_xlsx(data: bytes) -> ParsedDocument:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        pages: List[ParsedPage] = []
        for sheet_num, sheet in enumerate(wb.worksheets, start=1):
            rows: List[str] = []
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                row_text = " | ".join(cells).strip(" |")
                if row_text:
                    rows.append(row_text)
            if rows:
                pages.append(ParsedPage(page_number=sheet_num, text="\n".join(rows), metadata={"sheet_name": sheet.title}))
        wb.close()
        return ParsedDocument(pages=pages)
    except Exception as e:
        logger.error("XLSX parse failed: %s", e)
        return ParsedDocument(pages=[])


def parse_pptx(data: bytes) -> ParsedDocument:
    try:
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        pages: List[ParsedPage] = []
        for slide_num, slide in enumerate(prs.slides, start=1):
            texts: List[str] = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = "".join(run.text for run in para.runs)
                        if text.strip():
                            texts.append(text.strip())
            if texts:
                pages.append(ParsedPage(page_number=slide_num, text="\n".join(texts), metadata={"slide": slide_num}))
        return ParsedDocument(pages=pages)
    except Exception as e:
        logger.error("PPTX parse failed: %s", e)
        return ParsedDocument(pages=[])


def parse_html(data: bytes, encoding: str = "utf-8") -> ParsedDocument:
    try:
        from bs4 import BeautifulSoup
        try:
            text_raw = data.decode(encoding)
        except UnicodeDecodeError:
            import chardet
            detected = chardet.detect(data)
            text_raw = data.decode(detected.get("encoding", "utf-8"), errors="replace")
        soup = BeautifulSoup(text_raw, "lxml")
        for tag in soup(["script", "style", "nav", "footer", "header", "aside"]):
            tag.decompose()
        title = soup.title.string.strip() if soup.title else None
        body_text = soup.get_text(separator="\n", strip=True)
        pages = [ParsedPage(page_number=1, text=body_text)] if body_text.strip() else []
        return ParsedDocument(pages=pages, title=title)
    except Exception as e:
        logger.error("HTML parse failed: %s", e)
        return ParsedDocument(pages=[])


def parse_markdown(data: bytes) -> ParsedDocument:
    try:
        import markdown
        from bs4 import BeautifulSoup
        text_raw = data.decode("utf-8", errors="replace")
        html = markdown.markdown(text_raw)
        soup = BeautifulSoup(html, "lxml")
        plain = soup.get_text(separator="\n", strip=True)
        pages = [ParsedPage(page_number=1, text=plain)] if plain.strip() else []
        return ParsedDocument(pages=pages)
    except Exception as e:
        logger.error("Markdown parse failed: %s", e)
        text = data.decode("utf-8", errors="replace")
        return ParsedDocument(pages=[ParsedPage(page_number=1, text=text)])


def parse_text(data: bytes) -> ParsedDocument:
    try:
        import chardet
        detected = chardet.detect(data)
        encoding = detected.get("encoding") or "utf-8"
        text = data.decode(encoding, errors="replace")
        pages = [ParsedPage(page_number=1, text=text)] if text.strip() else []
        return ParsedDocument(pages=pages)
    except Exception as e:
        logger.error("Text parse failed: %s", e)
        return ParsedDocument(pages=[])


MIME_PARSERS = {
    "application/pdf": parse_pdf,
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": parse_docx,
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": parse_xlsx,
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": parse_pptx,
    "text/html": parse_html,
    "text/markdown": parse_markdown,
    "text/plain": parse_text,
    "text/csv": parse_text,
}

EXT_PARSERS = {
    ".pdf": parse_pdf,
    ".docx": parse_docx, ".doc": parse_docx,
    ".xlsx": parse_xlsx, ".xls": parse_xlsx,
    ".pptx": parse_pptx, ".ppt": parse_pptx,
    ".html": parse_html, ".htm": parse_html,
    ".md": parse_markdown, ".markdown": parse_markdown,
    ".txt": parse_text, ".csv": parse_text,
    ".rtf": parse_text, ".json": parse_text, ".xml": parse_text,
}


def parse_document(data: bytes, filename: str, mime_type: Optional[str] = None) -> ParsedDocument:
    if mime_type and mime_type in MIME_PARSERS:
        logger.info("Parsing %s via mime_type=%s", filename, mime_type)
        return MIME_PARSERS[mime_type](data)
    ext = Path(filename).suffix.lower()
    if ext in EXT_PARSERS:
        logger.info("Parsing %s via extension=%s", filename, ext)
        return EXT_PARSERS[ext](data)
    try:
        import magic
        detected_mime = magic.from_buffer(data[:2048], mime=True)
        if detected_mime in MIME_PARSERS:
            return MIME_PARSERS[detected_mime](data)
    except Exception:
        pass
    logger.warning("Unknown format for %s, trying plain text", filename)
    return parse_text(data)
